"""V2 工作底稿在线编辑 + 状态机 + 报告读底稿测试。"""
import json
from io import BytesIO

import pytest
from fastapi.testclient import TestClient


@pytest.fixture(scope="module")
def client(auth_headers):
    from app.main import app
    from app.seeds.load_indicators_55 import load
    load(replace=True)
    with TestClient(app, headers=auth_headers) as c:
        yield c


def _setup_task_with_worksheet(client, name):
    """建一个走到 ai_done 的任务，自动生成 worksheet。"""
    uid = client.post("/api/units", json={"name": name, "code": ""}).json()["id"]
    inds = client.get("/api/indicators").json()
    task_id = client.post("/api/tasks", json={
        "unit_id": uid, "name": name, "eval_year": 2025,
        "scope": "selected", "selected_indicator_ids": [inds[0]["id"], inds[1]["id"]],
    }).json()["id"]
    client.post(f"/api/tasks/{task_id}/materials",
                files={"file": ("a.txt", b"x", "text/plain")})
    client.post(f"/api/tasks/{task_id}/run")
    return task_id


def test_patch_audited_score(client):
    task_id = _setup_task_with_worksheet(client, "V2-EDIT-1")
    ws = client.get(f"/api/tasks/{task_id}/worksheet").json()
    row = ws["rows"][0]
    rid = row["id"]
    # 取一个与现值不同的合法值
    orig = float(row["audited_score"])
    new_value = 0.5 if orig != 0.5 else 1.0
    r = client.patch(f"/api/tasks/{task_id}/worksheet/rows/{rid}",
                     json={"audited_score": new_value})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["audited_score"] == new_value
    # 状态从 draft 升到 reviewing（有实际变更）
    assert body["worksheet_status"] == "reviewing"


def test_patch_audit_finding_text(client):
    task_id = _setup_task_with_worksheet(client, "V2-EDIT-2")
    ws = client.get(f"/api/tasks/{task_id}/worksheet").json()
    rid = ws["rows"][0]["id"]
    new_text = "审计师专业评语：制度齐全，建议补充绩效评价章节。"
    r = client.patch(f"/api/tasks/{task_id}/worksheet/rows/{rid}",
                     json={"audit_finding_text": new_text})
    assert r.status_code == 200
    assert r.json()["audit_finding_text"] == new_text


def test_patch_material_flags(client):
    task_id = _setup_task_with_worksheet(client, "V2-EDIT-3")
    ws = client.get(f"/api/tasks/{task_id}/worksheet").json()
    rid = ws["rows"][0]["id"]
    flags = {"real": False, "fake": True, "relevant": True, "irrelevant": False,
             "effective": True, "ineffective": False, "complete": True, "incomplete": False,
             "compliant": True, "non_compliant": False, "duplicate": False, "unique": True,
             "match_high": True, "match_low": False}
    r = client.patch(f"/api/tasks/{task_id}/worksheet/rows/{rid}",
                     json={"material_flags": flags})
    assert r.status_code == 200
    out_flags = json.loads(r.json()["material_flags"])
    assert out_flags["fake"] is True
    assert out_flags["real"] is False


def test_audited_score_out_of_range_rejected(client):
    task_id = _setup_task_with_worksheet(client, "V2-EDIT-4")
    ws = client.get(f"/api/tasks/{task_id}/worksheet").json()
    rid = ws["rows"][0]["id"]
    r = client.patch(f"/api/tasks/{task_id}/worksheet/rows/{rid}",
                     json={"audited_score": 9999})
    assert r.status_code == 400


def test_finalize_and_lock(client):
    task_id = _setup_task_with_worksheet(client, "V2-FINAL")
    # 定稿
    r = client.post(f"/api/tasks/{task_id}/worksheet/finalize")
    assert r.status_code == 200
    assert r.json()["status"] == "finalized"

    # finalized 后再改任何字段 → 400
    ws = client.get(f"/api/tasks/{task_id}/worksheet").json()
    rid = ws["rows"][0]["id"]
    r = client.patch(f"/api/tasks/{task_id}/worksheet/rows/{rid}",
                     json={"audited_score": 0.5})
    assert r.status_code == 400
    assert "定稿" in r.json()["detail"]

    # rebuild 也拒绝
    r = client.post(f"/api/tasks/{task_id}/worksheet/rebuild")
    assert r.status_code == 400


def test_unlock_requires_super_admin(client):
    task_id = _setup_task_with_worksheet(client, "V2-UNLOCK")
    client.post(f"/api/tasks/{task_id}/worksheet/finalize")
    # 当前 client 是 admin (super_admin) → 应该可以
    r = client.post(f"/api/tasks/{task_id}/worksheet/unlock")
    assert r.status_code == 200
    assert r.json()["status"] == "reviewing"


def test_report_uses_worksheet_text(client):
    """报告生成时应包含底稿里编辑过的核查情况说明。"""
    task_id = _setup_task_with_worksheet(client, "V2-REPORT")
    ws = client.get(f"/api/tasks/{task_id}/worksheet").json()
    row = ws["rows"][0]
    rid = row["id"]
    # v1.7 后：orchestrator 不再为无材料指标自动写 finding，
    # 而报告的"各指标核查明细"只迭代有 finding 的指标，
    # 因此显式注入一条 finding 以便 audit_finding_text 被渲染。
    from app.models import Finding, SessionLocal
    with SessionLocal() as s:
        s.add(Finding(task_id=task_id, indicator_id=row["indicator_id"],
                      finding_type="合规性问题", severity="低",
                      description="V2-REPORT 测试用 finding",
                      review_status="pending", source="rule"))
        s.commit()
    custom_text = "审计师独家批注：制度齐全且公章清晰，建议保持原分。"
    client.patch(f"/api/tasks/{task_id}/worksheet/rows/{rid}",
                 json={"audit_finding_text": custom_text})

    # 下载 docx 并检查内容
    r = client.get(f"/api/tasks/{task_id}/report")
    assert r.status_code == 200
    from docx import Document
    doc = Document(BytesIO(r.content))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    # 报告应包含底稿编辑的文字
    assert "审计师独家批注" in full_text, "Word 报告未包含底稿审计师评语"
